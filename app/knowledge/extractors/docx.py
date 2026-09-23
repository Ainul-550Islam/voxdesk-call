"""
DOCX extraction preserving paragraph order and heading structure.

python-docx reads the XML directly; it never opens Word and never runs a
macro. A .docx is a zip, so a corrupt archive is the common failure and is
reported as such rather than escaping as a raw zipfile error.
"""
from __future__ import annotations

import io

import structlog

from app.knowledge.extractors.base import (
    ExtractedDocument,
    ExtractionError,
    Section,
    clean_text,
)

log = structlog.get_logger()


class DocxExtractor:
    def extract(self, data: bytes, *, filename: str | None = None) -> ExtractedDocument:
        if not data:
            raise ExtractionError("the file is empty")

        try:
            import docx
        except ImportError as exc:      # pragma: no cover - deployment issue
            raise ExtractionError("DOCX support is not installed") from exc

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            log.warning("extract.docx_unreadable", error=str(exc))
            raise ExtractionError("this Word document could not be opened") from exc

        sections: list[Section] = []
        current_heading: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            body = clean_text("\n".join(buffer))
            if body:
                meta = {"heading": current_heading} if current_heading else {}
                sections.append(Section(text=body, metadata=meta))
            buffer.clear()

        # Paragraph order is the document's own reading order, so it is
        # preserved exactly rather than regrouped.
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            style = (paragraph.style.name if paragraph.style else "") or ""

            if style.startswith("Heading") and text:
                flush()
                current_heading = text
                buffer.append(text)
                continue
            if text:
                buffer.append(text)
        flush()

        # Tables carry the price lists people actually ask about, so they are
        # flattened into readable rows rather than dropped.
        table_rows: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))
        if table_rows:
            sections.append(
                Section(
                    text=clean_text("\n".join(table_rows)),
                    metadata={"block": "tables"},
                )
            )

        if not sections:
            raise ExtractionError("this Word document contains no readable text")

        return ExtractedDocument(
            text="\n\n".join(s.text for s in sections),
            sections=sections,
            metadata={
                "format": "docx",
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
        )